"""
Document Text and Metadata Extractor
Supports PDF (pypdf), DOCX (python-docx), and TXT files.
"""

import io
import os
from typing import Dict, Any, Tuple
import pypdf
import docx


def extract_text_and_metadata(filename: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts raw text and metadata from uploaded file bytes.
    """
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    metadata: Dict[str, Any] = {
        "filename": filename,
        "extension": ext,
        "file_size_bytes": len(file_bytes),
        "pages_or_sections": 1,
        "author": "Unknown",
        "creator": "Unknown",
        "producer": "Unknown",
        "created_date": None,
        "modified_date": None,
    }

    if ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            metadata["pages_or_sections"] = len(reader.pages)
            pages_text = []
            for page in reader.pages:
                page_content = page.extract_text() or ""
                pages_text.append(page_content)
            text = "\n".join(pages_text)

            pdf_meta = reader.metadata
            if pdf_meta:
                metadata["author"] = str(pdf_meta.get("/Author") or "Not Specified")
                metadata["creator"] = str(pdf_meta.get("/Creator") or "Not Specified")
                metadata["producer"] = str(pdf_meta.get("/Producer") or "Not Specified")
                metadata["created_date"] = str(pdf_meta.get("/CreationDate") or "")
                metadata["modified_date"] = str(pdf_meta.get("/ModDate") or "")
        except Exception as e:
            text = f"[PDF Extraction Error: {str(e)}]"

    elif ext in [".docx", ".doc"]:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n".join(paragraphs)
            metadata["pages_or_sections"] = max(1, len(doc.paragraphs) // 15)

            props = doc.core_properties
            if props:
                metadata["author"] = props.author or "Not Specified"
                metadata["created_date"] = str(props.created) if props.created else ""
                metadata["modified_date"] = str(props.modified) if props.modified else ""
                metadata["revision"] = props.revision
        except Exception as e:
            text = f"[DOCX Extraction Error: {str(e)}]"

    else:
        # Assume plain text / UTF-8
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1")
            except Exception:
                text = str(file_bytes)
        metadata["pages_or_sections"] = max(1, len(text.split("\n\n")))

    return {
        "text": text.strip(),
        "metadata": metadata
    }
